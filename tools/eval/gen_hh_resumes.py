"""Генератор обезличенного набора резюме в формате выгрузки hh.ru.

    uv run python tools/eval/gen_hh_resumes.py   # -> data/eval/resumes/hh/

Люди, компании, телефоны и почты вымышленные (адреса `example.test`, телефоны вида 900-000-xx-xx).
Раскладка повторяет выгрузку hh.ru: шапка, «Желаемая должность и зарплата», «Опыт работы N лет»,
слева даты, справа компания, город и сайт, должность и обязанности; «Образование» с годом
отдельной строкой; «Ключевые навыки» → «Знание языков» и «Навыки»; «Обо мне». Нечётные резюме —
PDF, чётные — DOCX (опыт таблицей, как в выгрузке Word). Эталон `gold.json` — то, что должен
извлечь разбор (`tools/eval/runners/resumes.py`).
"""

from __future__ import annotations

import json
import math
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "eval" / "resumes" / "hh"
TODAY = (2026, 9)
MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь",
          "Октябрь", "Ноябрь", "Декабрь"]  # fmt: skip

# exp: (начало, конец|None, компания, сайт|None, город, должность, обязанности)
# edu: (уровень, год, вуз, город, факультет, специальность)
PEOPLE: list[dict] = [
    dict(
        name="Воронов Артём Игоревич",
        sex="Мужчина, 29 лет",
        phone="+7 (900) 000-11-01",
        email="a.voronov@example.test",
        city="Москва",
        wish="Backend-разработчик",
        salary="250 000 ₽ на руки",
        exp=[
            (
                (2021, 8),
                None,
                "ООО «Вымышленная логистика»",
                "vlogistics.example.test",
                "Москва",
                "Backend-разработчик",
                "Делал сервисы маршрутизации на Python и FastAPI. Хранил данные в PostgreSQL, "
                "кеш в Redis. Настроил CI в GitLab.",
            ),
            (
                (2018, 6),
                (2021, 7),
                "АО «Придуманные системы»",
                None,
                "Москва",
                "Python-разработчик",
                "Поддерживал внутренний портал на Django, писал тесты на pytest.",
            ),
        ],
        edu=[
            (
                "Высшее образование",
                2018,
                "Московский вымышленный университет",
                "Москва",
                "Факультет вычислительной математики",
                "Прикладная математика и информатика",
            )
        ],
        langs=["Русский — Родной", "Английский — B2 — Средне-продвинутый"],
        skills=["Python", "FastAPI", "Django", "PostgreSQL", "Redis", "Docker", "Git", "pytest"],
        about="Пишу надёжные сервисы и люблю разбирать инциденты. Ищу команду с сильной инженерной "
        "культурой.",
    ),
    dict(
        name="Лебедева Ксения Андреевна",
        sex="Женщина, 26 лет",
        phone="+7 (900) 000-11-02",
        email="k.lebedeva@example.test",
        city="Санкт-Петербург",
        wish="Data Scientist",
        salary="220 000 ₽ на руки",
        exp=[
            (
                (2022, 9),
                None,
                "ООО «Придуманные данные»",
                "pdata.example.test",
                "Санкт-Петербург",
                "Data Scientist",
                "Строила модели ранжирования на PyTorch и scikit-learn, данные готовила в pandas и "
                "Airflow.",
            ),
            (
                (2020, 2),
                (2022, 8),
                "ООО «Вымышленный банк»",
                None,
                "Санкт-Петербург",
                "Аналитик данных",
                "Отчёты на SQL и Python, дашборды в Grafana.",
            ),
        ],
        edu=[
            (
                "Магистр",
                2020,
                "Санкт-Петербургский вымышленный университет",
                "Санкт-Петербург",
                "Факультет информатики",
                "Машинное обучение",
            )
        ],
        langs=["Русский — Родной", "Английский — C1 — Продвинутый"],
        skills=["Python", "PyTorch", "scikit-learn", "pandas", "SQL", "Airflow", "Grafana"],
        about="Занимаюсь рекомендательными системами. Люблю воспроизводимые эксперименты.",
    ),
    dict(
        name="Соколов Дмитрий Павлович",
        sex="Мужчина, 33 года",
        phone="+7 (900) 000-11-03",
        email="d.sokolov@example.test",
        city="Казань",
        wish="Frontend-разработчик",
        salary="200 000 ₽ на руки",
        exp=[
            (
                (2020, 3),
                (2026, 5),
                "ООО «Выдуманные интерфейсы»",
                "vui.example.test",
                "Казань",
                "Ведущий frontend-разработчик",
                "Вёл команду из четырёх человек, писал интерфейсы на React и TypeScript, внедрил "
                "дизайн-систему.",
            ),
            (
                (2016, 9),
                (2020, 2),
                "ИП Тестов",
                None,
                "Казань",
                "Верстальщик",
                "Верстал сайты на HTML и CSS, подключал JavaScript.",
            ),
        ],
        edu=[
            (
                "Высшее образование",
                2016,
                "Казанский вымышленный технический университет",
                "Казань",
                "Факультет информационных технологий",
                "Программная инженерия",
            )
        ],
        langs=["Русский — Родной", "Английский — B1 — Средний"],
        skills=["React", "TypeScript", "JavaScript", "HTML", "CSS", "Git", "Figma"],
        about=None,
    ),
    dict(
        name="Кузнецова Анна Викторовна",
        sex="Женщина, 31 год",
        phone="+7 (900) 000-11-04",
        email="a.kuznetsova@example.test",
        city="Новосибирск",
        wish="DevOps-инженер",
        salary="260 000 ₽ на руки",
        exp=[
            (
                (2023, 1),
                None,
                "ООО «Придуманный хостинг»",
                "phost.example.test",
                "Новосибирск",
                "DevOps-инженер",
                "Поддерживала кластеры Kubernetes, описывала инфраструктуру в Terraform и Ansible, "
                "мониторинг в Grafana.",
            ),
            (
                (2019, 4),
                (2022, 12),
                "АО «Вымышленные сети»",
                None,
                "Новосибирск",
                "Системный администратор",
                "Администрировала серверы на Linux, писала скрипты на Bash.",
            ),
        ],
        edu=[
            (
                "Высшее образование",
                2017,
                "Новосибирский вымышленный государственный университет",
                "Новосибирск",
                "Факультет информационных технологий",
                "Информационные системы",
            )
        ],
        langs=["Русский — Родной", "Английский — B2 — Средне-продвинутый"],
        skills=["Kubernetes", "Docker", "Terraform", "Ansible", "Linux", "Grafana", "Bash"],
        about="Автоматизирую всё, что делается руками больше двух раз.",
    ),
    dict(
        name="Морозов Илья Сергеевич",
        sex="Мужчина, 24 года",
        phone="+7 (900) 000-11-05",
        email="i.morozov@example.test",
        city="Екатеринбург",
        wish="QA-инженер",
        salary="150 000 ₽ на руки",
        exp=[
            (
                (2024, 6),
                None,
                "ООО «Выдуманный софт»",
                "vsoft.example.test",
                "Екатеринбург",
                "Инженер по тестированию",
                "Автотесты на Python и Selenium, ручное тестирование, ведение дефектов в Jira.",
            ),
            (
                (2023, 2),
                (2024, 5),
                "ООО «Придуманные игры»",
                None,
                "Екатеринбург",
                "Стажёр QA",
                "Проверял сборки, писал тест-кейсы.",
            ),
        ],
        edu=[
            (
                "Бакалавр",
                2024,
                "Уральский вымышленный федеральный университет",
                "Екатеринбург",
                "Институт радиоэлектроники и информационных технологий",
                "Информатика",
            )
        ],
        langs=["Русский — Родной", "Английский — B1 — Средний"],
        skills=["Python", "Selenium", "pytest", "SQL", "Git", "Jira"],
        about="Хочу вырасти до автоматизатора тестирования.",
    ),
    dict(
        name="Павлова Елена Николаевна",
        sex="Женщина, 28 лет",
        phone="+7 (900) 000-11-06",
        email="e.pavlova@example.test",
        city="Москва",
        wish="Android-разработчик",
        salary="230 000 ₽ на руки",
        exp=[
            (
                (2021, 11),
                None,
                "ООО «Вымышленный маркет»",
                "vmarket.example.test",
                "Москва",
                "Android-разработчик",
                "Развивала приложение на Kotlin, подключала платежи, писала тесты.",
            ),
            (
                (2019, 9),
                (2021, 10),
                "ООО «Придуманные приложения»",
                None,
                "Москва",
                "Junior Android-разработчик",
                "Верстала экраны, исправляла ошибки, работала с Git.",
            ),
        ],
        edu=[
            (
                "Высшее образование",
                2019,
                "Московский вымышленный технический институт",
                "Москва",
                "Факультет прикладной информатики",
                "Программная инженерия",
            )
        ],
        langs=["Русский — Родной", "Английский — B2 — Средне-продвинутый"],
        skills=["Kotlin", "Git", "SQL", "Docker", "Jira"],
        about="Люблю аккуратные интерфейсы и быстрые приложения.",
    ),
    dict(
        name="Николаев Григорий Олегович",
        sex="Мужчина, 36 лет",
        phone="+7 (900) 000-11-07",
        email="g.nikolaev@example.test",
        city="Самара",
        wish="Инженер данных",
        salary="300 000 ₽ на руки",
        exp=[
            (
                (2019, 5),
                None,
                "ПАО «Вымышленная энергия»",
                "venergy.example.test",
                "Самара",
                "Ведущий инженер данных",
                "Построил хранилище на ClickHouse, потоки на Kafka и Airflow, перенёс отчётность.",
            ),
            (
                (2014, 8),
                (2019, 4),
                "ООО «Придуманные данные»",
                None,
                "Самара",
                "Разработчик SQL",
                "Писал хранимые процедуры, оптимизировал запросы к PostgreSQL.",
            ),
        ],
        edu=[
            (
                "Высшее образование",
                2012,
                "Самарский вымышленный университет",
                "Самара",
                "Факультет информатики",
                "Прикладная информатика",
            )
        ],
        langs=["Русский — Родной", "Английский — B1 — Средний"],
        skills=["ClickHouse", "Kafka", "Airflow", "PostgreSQL", "SQL", "Python", "Linux"],
        about="Строю надёжные конвейеры данных.",
    ),
    dict(
        name="Захарова Ольга Дмитриевна",
        sex="Женщина, 30 лет",
        phone="+7 (900) 000-11-08",
        email="o.zakharova@example.test",
        city="Нижний Новгород",
        wish="Системный аналитик",
        salary="180 000 ₽ на руки",
        exp=[
            (
                (2022, 2),
                None,
                "ООО «Выдуманные решения»",
                "vsol.example.test",
                "Нижний Новгород",
                "Системный аналитик",
                "Описывала требования, проектировала API, вела документацию в Confluence и Jira.",
            ),
            (
                (2018, 10),
                (2022, 1),
                "АО «Придуманный банк»",
                None,
                "Нижний Новгород",
                "Бизнес-аналитик",
                "Собирала требования, строила отчёты в Excel и SQL.",
            ),
        ],
        edu=[
            (
                "Высшее образование",
                2018,
                "Нижегородский вымышленный университет",
                "Нижний Новгород",
                "Факультет экономики",
                "Бизнес-информатика",
            )
        ],
        langs=["Русский — Родной", "Английский — B2 — Средне-продвинутый"],
        skills=["SQL", "Excel", "Jira", "Scrum", "Python"],
        about="Перевожу пожелания бизнеса на язык разработки.",
    ),
    dict(
        name="Беляев Максим Александрович",
        sex="Мужчина, 27 лет",
        phone="+7 (900) 000-11-09",
        email="m.belyaev@example.test",
        city="Воронеж",
        wish="Fullstack-разработчик",
        salary="210 000 ₽ на руки",
        exp=[
            (
                (2023, 4),
                (2026, 8),
                "ООО «Вымышленные сервисы»",
                "vserv.example.test",
                "Воронеж",
                "Fullstack-разработчик",
                "Делал интерфейсы на Vue и серверную часть на Python и Flask, деплой через Docker.",
            ),
            (
                (2020, 1),
                (2023, 3),
                "ИП Тестов",
                None,
                "Воронеж",
                "Веб-разработчик",
                "Разрабатывал сайты на JavaScript и PHP.",
            ),
        ],
        edu=[
            (
                "Среднее специальное образование",
                2019,
                "Воронежский вымышленный колледж",
                "Воронеж",
                "Отделение информационных технологий",
                "Программирование в компьютерных системах",
            )
        ],
        langs=["Русский — Родной"],
        skills=["Vue", "JavaScript", "Python", "Flask", "Docker", "PostgreSQL", "Git"],
        about=None,
    ),
    dict(
        name="Тихонова Дарья Максимовна",
        sex="Женщина, 25 лет",
        phone="+7 (900) 000-11-10",
        email="d.tikhonova@example.test",
        city="Томск",
        wish="ML-инженер",
        salary="240 000 ₽ на руки",
        exp=[
            (
                (2023, 7),
                None,
                "ООО «Придуманный интеллект»",
                "pai.example.test",
                "Томск",
                "ML-инженер",
                "Обучала нейросети на PyTorch, готовила данные в NumPy и pandas, внедряла модели "
                "в Docker.",
            ),
        ],
        edu=[
            (
                "Магистр",
                2023,
                "Томский вымышленный политехнический университет",
                "Томск",
                "Институт кибернетики",
                "Искусственный интеллект",
            ),
            (
                "Бакалавр",
                2021,
                "Томский вымышленный политехнический университет",
                "Томск",
                "Институт кибернетики",
                "Прикладная информатика",
            ),
        ],
        langs=["Русский — Родной", "Английский — C1 — Продвинутый"],
        skills=["Python", "PyTorch", "NumPy", "pandas", "Docker", "Linux", "Git"],
        about="Интересуюсь компьютерным зрением и оптимизацией моделей.",
    ),
]


def _plural(k: int, one: str, few: str, many: str) -> str:
    k10, k100 = k % 10, k % 100
    if k10 == 1 and k100 != 11:
        return one
    return few if 2 <= k10 <= 4 and not 12 <= k100 <= 14 else many


def _span(n: int) -> str:
    y, m = divmod(n, 12)
    parts = []
    if y:
        parts.append(f"{y} {_plural(y, 'год', 'года', 'лет')}")
    if m:
        parts.append(f"{m} {_plural(m, 'месяц', 'месяца', 'месяцев')}")
    return " ".join(parts) or "1 месяц"


def _n(a: tuple[int, int], b: tuple[int, int] | None) -> int:
    end = b or TODAY
    return (end[0] - a[0]) * 12 + end[1] - a[1] + 1


def _date(d: tuple[int, int]) -> str:
    return f"{MONTHS[d[1] - 1]} {d[0]}"


def _ym(d: tuple[int, int] | None) -> str | None:
    return None if d is None else f"{d[0]:04d}-{d[1]:02d}"


def build(p: dict) -> list[tuple]:
    """Блоки: (вид, текст) или ('exp', период, длительность, правые строки)."""
    total = sum(_n(e[0], e[1]) for e in p["exp"])
    b: list[tuple] = [
        ("title", p["name"]),
        ("line", p["sex"]),
        ("line", f"{p['phone']} — предпочитаемый способ связи"),
        ("line", p["email"]),
        ("line", f"Проживает: {p['city']}"),
        ("line", "Гражданство: Россия, есть разрешение на работу: Россия"),
        ("line", "Не готов к переезду, не готов к командировкам"),
        ("head", "Желаемая должность и зарплата"),
        ("line", p["wish"]),
        ("line", "Специализации:"),
        ("line", "— Программист, разработчик"),
        ("line", "Занятость: полная занятость"),
        ("line", "График работы: полный день"),
        ("line", p["salary"]),
        ("head", f"Опыт работы — {_span(total)}"),
    ]
    for start, end, company, site, city, title, duty in p["exp"]:
        place = f"{city}, {site}" if site else city
        period = f"{_date(start)} — {_date(end) if end else 'настоящее время'}"
        b.append(
            ("exp", period, _span(_n(start, end)), [company, place, title, f"Обязанности: {duty}"])
        )
    b.append(("head", "Образование"))
    level = None
    for lv, year, inst, city, faculty, spec in p["edu"]:
        if lv != level:
            b.append(("line", lv))
            level = lv
        b += [("line", str(year)), ("line", f"{inst}, {city}"), ("line", faculty), ("line", spec)]
    b += [("head", "Ключевые навыки"), ("sub", "Знание языков")]
    b += [("line", ln) for ln in p["langs"]]
    b += [("sub", "Навыки"), ("line", "  ".join(p["skills"]))]
    if p["about"]:
        b += [("head", "Обо мне"), ("line", p["about"])]
    b.append(("line", "Резюме обновлено 12 сентября 2026 в 12:30"))
    return b


def gold(p: dict, skill_slug) -> dict:
    return {
        "full_name": p["name"],
        "phone": "+" + re.sub(r"\D", "", p["phone"]),
        "email": p["email"],
        "city": p["city"],
        "summary": p["about"],
        "experience": [
            {"company": e[2], "title": e[5], "start": _ym(e[0]), "end": _ym(e[1])} for e in p["exp"]
        ],
        "education": [{"institution": e[2], "year_end": e[1], "field": e[5]} for e in p["edu"]],
        "skills": sorted({s for n in p["skills"] if (s := skill_slug(n))}),
    }


def write_pdf(blocks: list[tuple], path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    st: dict = {"page": None, "y": 0.0}

    def put(x: float, w: float, html: str, h: float) -> None:
        if st["page"] is None or st["y"] + h > 800:
            st["page"] = doc.new_page()
            st["y"] = 40.0
        st["page"].insert_htmlbox(pymupdf.Rect(x, st["y"], x + w, st["y"] + h), html)

    for blk in blocks:
        kind = blk[0]
        if kind == "title":
            put(36, 520, f"<b style='font-size:17pt'>{blk[1]}</b>", 26)
            st["y"] += 26
        elif kind in ("head", "sub"):
            st["y"] += 8 if kind == "head" else 2
            size = 13 if kind == "head" else 11
            put(36, 520, f"<b style='font-size:{size}pt'>{blk[1]}</b>", 20)
            st["y"] += 20
        elif kind == "line":
            h = 14 * max(1, math.ceil(len(blk[1]) / 85))
            put(36, 520, f"<span style='font-size:10pt'>{blk[1]}</span>", h)
            st["y"] += h
        else:
            _, period, duration, right = blk
            y0 = st["y"]
            for i, r in enumerate(right):
                h = 14 * max(1, math.ceil(len(r) / 62))
                text = f"<b>{r}</b>" if i == 0 else r
                put(210, 350, f"<span style='font-size:10pt'>{text}</span>", h)
                st["y"] += h
            y_end = st["y"]
            st["y"] = y0
            put(36, 165, f"<b style='font-size:10pt'>{period}</b>", 14)
            st["y"] += 14
            put(36, 165, f"<span style='font-size:10pt'>{duration}</span>", 14)
            st["y"] = y_end + 8
    path.write_bytes(doc.tobytes())


def write_docx(blocks: list[tuple], path: Path) -> None:
    import docx

    d = docx.Document()
    for blk in blocks:
        kind = blk[0]
        if kind == "title":
            d.add_heading(blk[1], level=1)
        elif kind == "head":
            d.add_heading(blk[1], level=2)
        elif kind == "sub":
            d.add_heading(blk[1], level=3)
        elif kind == "line":
            d.add_paragraph(blk[1])
        else:
            _, period, duration, right = blk
            left, rt = d.add_table(rows=1, cols=2).rows[0].cells
            left.text = period
            left.add_paragraph(duration)
            rt.text = right[0]
            for r in right[1:]:
                rt.add_paragraph(r)
    d.save(path)


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    import yaml
    from fsp_ml.taxonomy import TaxonomyIndex

    skills = yaml.safe_load((ROOT / "data/taxonomy/skills.yaml").read_text("utf-8"))["skills"]
    index = TaxonomyIndex.from_snapshot({"version": 1, "skills": skills})
    OUT.mkdir(parents=True, exist_ok=True)
    items = []
    for i, p in enumerate(PEOPLE, 1):
        name = f"{i:02d}." + ("pdf" if i % 2 else "docx")
        (write_pdf if i % 2 else write_docx)(build(p), OUT / name)
        items.append({"file": name, **gold(p, index.lookup)})
    (OUT / "gold.json").write_text(
        json.dumps({"version": 1, "resumes": items}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"{len(items)} резюме → {OUT}")


if __name__ == "__main__":
    main()
