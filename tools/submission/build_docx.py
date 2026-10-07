"""AU-4: документ сдачи Markdown → DOCX (ТЗ: «doc или pdf»).

Зависимость уже есть в workspace (`python-docx` у services/ml), новых нет. Разбор Markdown
общий с PDF-сборкой (`build_pdf`): заголовки, абзацы, вложенные списки, таблицы, код, цитаты,
**жирный**, *курсив*, `код`, ссылки (в DOCX — текст ссылки, адрес не печатаем).
Запуск: `uv run python tools/submission/build_docx.py [--src docs/submission.md] [--out reports/submission/submission.docx]`.
"""

from __future__ import annotations

import argparse
import importlib.util
import re
import sys
from pathlib import Path
from typing import Any

# ruff: noqa: E501
_HERE = Path(__file__).resolve().parent


def _load_pdf_helpers() -> Any:
    spec = importlib.util.spec_from_file_location("sub_build_pdf", _HERE / "build_pdf.py")
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


_pdf = _load_pdf_helpers()

_TOKEN = re.compile(
    r"`([^`]+)`|\*\*(.+?)\*\*|(?<![\w*])\*([^*\s][^*]*?)\*(?![\w*])|\[([^\]]+)\]\(([^)\s]+)\)"
)


def add_runs(par: Any, text: str, bold: bool = False) -> None:
    """Раскладывает inline-разметку по ранам параграфа."""
    pos = 0
    for m in _TOKEN.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos : m.start()]).bold = bold or None
        if m.group(1) is not None:
            run = par.add_run(m.group(1))
            run.font.name = "Consolas"
            run.bold = bold or None
        elif m.group(2) is not None:
            add_runs(par, m.group(2), bold=True)
        elif m.group(3) is not None:
            run = par.add_run(m.group(3))
            run.italic = True
            run.bold = bold or None
        else:
            par.add_run(m.group(4)).bold = bold or None
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:]).bold = bold or None


def build_document(md: str) -> Any:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    lines = md.splitlines()
    i = 0
    para: list[str] = []

    def flush() -> None:
        if para:
            add_runs(doc.add_paragraph(), " ".join(para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            flush()
            lang = line[3:].strip()
            i += 1
            buf: list[str] = []
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            if lang == "mermaid":
                note = doc.add_paragraph()
                note.add_run(
                    "Схема (Mermaid) отрисовывается в Markdown-версии; ниже её исходный текст."
                ).italic = True
            run = doc.add_paragraph().add_run("\n".join(buf))
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            continue
        if not line.strip():
            flush()
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            flush()
            heading = doc.add_heading(level=len(m.group(1)))
            add_runs(heading, m.group(2))
            i += 1
            continue
        if re.match(r"^-{3,}\s*$", line):
            flush()
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and _pdf._is_sep(lines[i + 1]):
            flush()
            head = _pdf._cells(line)
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(_pdf._cells(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(head))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, head, strict=False):
                add_runs(cell.paragraphs[0], text, bold=True)
            for r in rows:
                for cell, text in zip(table.add_row().cells, r, strict=False):
                    add_runs(cell.paragraphs[0], text)
            doc.add_paragraph()
            continue
        if line.startswith(">"):
            flush()
            buf = []
            while i < len(lines) and lines[i].startswith(">"):
                buf.append(lines[i].lstrip(">").strip())
                i += 1
            add_runs(doc.add_paragraph(style="Intense Quote"), " ".join(buf))
            continue
        lm = _pdf._LIST_RE.match(line)
        if lm:
            flush()
            base = "List Number" if lm.group(2)[0].isdigit() else "List Bullet"
            depth = min(len(lm.group(1)) // 2 + 1, 3)
            add_runs(
                doc.add_paragraph(style=base if depth == 1 else f"{base} {depth}"), lm.group(3)
            )
            i += 1
            continue
        para.append(line.strip())
        i += 1
    flush()
    return doc


def build(src: Path, out: Path) -> int:
    out.parent.mkdir(parents=True, exist_ok=True)
    build_document(src.read_text(encoding="utf-8")).save(out)
    return out.stat().st_size


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--src", type=Path, default=Path("docs/submission.md"))
    ap.add_argument("--out", type=Path, default=Path("reports/submission/submission.docx"))
    args = ap.parse_args()
    print(f"{args.out}: {build(args.src, args.out) // 1024} КБ")


if __name__ == "__main__":
    main()
