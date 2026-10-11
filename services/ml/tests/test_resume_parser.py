"""Разбор резюме (E3-3): текст, правила, LLM по схеме, проверка цитат, нормализация, прогресс.

LLM — фейковый клиент с записанным ответом (`fixtures/llm_resume_backend.json`), сеть не нужна.
"""

import copy
import io
import json
from datetime import date
from pathlib import Path
from uuid import uuid4

import httpx
import pytest
import yaml
from faststream.rabbit import RabbitQueue, TestRabbitBroker
from fsp_contracts import commands as cmd
from fsp_contracts import new_command
from fsp_contracts.ml import ResumeExperience
from fsp_ml import resume_parser, resume_rules
from fsp_ml.config import Settings
from fsp_ml.handlers import resume_parse_llm
from fsp_ml.llm import LlmError, OpenAiLlm
from fsp_ml.resume_text import DOCX_MIME, PDF_MIME, ResumeFileError, extract_text, sniff
from fsp_ml.taxonomy import TaxonomyIndex, TaxonomyProvider
from fsp_ml.worker import MlWorker
from fsp_platform.messaging import EVENTS_EXCHANGE, UI_EXCHANGE, durable_queue
from pydantic import ValidationError

FIX = Path(__file__).parent / "fixtures"
TAX = Path(__file__).resolve().parents[3] / "data" / "taxonomy"
TODAY = date(2026, 10, 1)
TEXT = (FIX / "resume_backend.txt").read_text()
LLM_OUT = json.loads((FIX / "llm_resume_backend.json").read_text())


def _index() -> TaxonomyIndex:
    skills = yaml.safe_load((TAX / "skills.yaml").read_text())["skills"]
    return TaxonomyIndex.from_snapshot({"version": 3, "skills": skills})


INDEX = _index()


class FakeLlm:
    """Записанные ответы по очереди; запоминает запросы."""

    def __init__(self, *answers):
        self.answers = list(answers)
        self.calls: list[dict] = []

    async def complete_json(self, messages, schema, *, schema_name, temperature=0.0):
        self.calls.append({"messages": messages, "schema": schema, "name": schema_name})
        a = self.answers.pop(0)
        if isinstance(a, Exception):
            raise a
        return copy.deepcopy(a)


def _pdf(text: str) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    html = "".join(f"<p>{line}</p>" for line in text.splitlines())
    page.insert_htmlbox(pymupdf.Rect(36, 36, 560, 800), html)
    return doc.tobytes()


def _docx(text: str, table: list[list[str]] | None = None) -> bytes:
    import docx

    d = docx.Document()
    for line in text.splitlines():
        d.add_paragraph(line)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, v in enumerate(row):
                t.cell(i, j).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# --- текст ---


def test_pdf_and_docx_text():
    pdf = _pdf(TEXT)
    assert sniff(pdf) == PDF_MIME
    got = extract_text(pdf)
    assert "Разрабатываю микросервисы на Python" in got and "Litestar" in got

    dx = _docx(TEXT, table=[["Навык", "Уровень"], ["Kafka", "базовый"]])
    assert sniff(dx) == DOCX_MIME
    got = extract_text(dx)
    assert "Postgres 15" in got and "Kafka | базовый" in got


def test_bad_and_scanned_files():
    with pytest.raises(ResumeFileError) as e:
        extract_text(b"MZ\x90\x00 not a resume" * 50)
    assert e.value.code == "bad_file"
    with pytest.raises(ResumeFileError) as e:
        extract_text(b"%PDF-1.7 broken")
    assert e.value.code == "bad_file"
    with pytest.raises(ResumeFileError) as e:
        extract_text(_pdf("Скан"))  # текстового слоя почти нет
    assert e.value.code == "scanned_pdf"


# --- правила ---


def test_contacts_by_rules():
    c = resume_rules.contacts(TEXT)
    assert c.email == "petr.vymyslov@example.test"
    assert c.phone == "+79001234567"
    assert c.telegram == "@petr_vymyslov"
    assert c.links == ["https://github.com/petr-vymyslov"]


@pytest.mark.parametrize(
    ("raw", "want"),
    [
        ("03.2020", "2020-03"),
        ("2022-03", "2022-03"),
        ("янв 2021", "2021-01"),
        ("Сентябрь 2019", "2019-09"),
        ("May 2018", "2018-05"),
        ("2019", "2019-01"),
        ("н.в.", None),
        ("когда-то", None),
        ("13.2020", None),
    ],
)
def test_parse_month(raw, want):
    assert resume_rules.parse_month(raw) == want


def test_parse_end():
    assert resume_rules.parse_end("н.в.") is None
    assert resume_rules.parse_end("по настоящее время") is None
    assert resume_rules.parse_end("2022") == "2022-12"
    assert resume_rules.parse_end("фев 2021") == "2021-02"


def test_experience_months_overlap_and_internship():
    exp = [
        ResumeExperience(title="Разработчик", start="2021-01", end="2022-01"),  # 12
        ResumeExperience(title="Разработчик", start="2021-07", end="2022-07"),  # +6 сверху
        ResumeExperience(title="Стажёр", start="2020-01", end="2020-07"),  # 6 × 0.5
        ResumeExperience(title="Пет-проект", start="2019-01", end="2020-01", is_commercial=False),
    ]
    assert resume_rules.total_experience_months(exp, TODAY) == 21
    # открытый период — до текущего месяца
    open_ = [ResumeExperience(title="Разработчик", start="2026-04", end=None)]
    assert resume_rules.total_experience_months(open_, TODAY) == 6


# --- проверка цитат ---


def test_quote_check():
    norm = resume_parser.norm_text(TEXT)
    assert resume_parser.quote_in_text("кеширование в   Redis.", norm)
    assert resume_parser.quote_in_text("Разрабатываю  микросервисы на python и FastAPI", norm)
    assert not resume_parser.quote_in_text("Разворачивал сервисы в Kubernetes", norm)
    q = resume_parser.norm_text("Немного писал на JavaScript для внутренних админок.")
    assert resume_parser.mention_in_quote("JavaScript", q, [])
    assert not resume_parser.mention_in_quote("Java", q, ["java se", "jdk"])  # Java ≠ JavaScript
    assert resume_parser.mention_in_quote("Python", "пишу на питоне", ["питон"])  # падеж
    assert not resume_parser.mention_in_quote("Go", "сервисы упаковываю в docker", ["golang"])


# --- сборка ---


def _assemble(llm_out=None, index=INDEX):
    out = resume_parser.LlmResume.model_validate(llm_out or LLM_OUT)
    return resume_parser.assemble(TEXT, out, index, version="resume-v1:test", today=TODAY)


def test_assemble_drops_skills_without_quote():
    p = _assemble()
    slugs = [s.slug for s in p.skills]
    assert slugs == [
        "python", "fastapi", "postgresql", "rabbitmq", "redis", "docker", "pytest", "git",
        "javascript",
    ]  # fmt: skip
    assert "kubernetes" not in slugs and "java" not in slugs and "go" not in slugs
    assert "quote_not_found:Kubernetes" in p.warnings
    assert "skill_not_in_quote:Java" in p.warnings
    assert "skill_not_in_quote:Go" in p.warnings
    assert p.unknown_skills == ["Litestar"]
    by = {s.slug: s for s in p.skills}
    assert by["postgresql"].raw == "Postgres 15" and by["postgresql"].method == "alias"
    assert all(0 < s.confidence <= 1 and s.quote for s in p.skills)


def test_assemble_experience_contacts_and_months():
    p = _assemble()
    work, intern = p.experience
    assert (work.start, work.end) == ("2021-03", None)
    assert (intern.start, intern.end) == ("2020-06", "2021-02")
    # навык опыта без подтверждённой цитаты (Kubernetes) не попадает
    assert work.skills == ["python", "fastapi", "postgresql", "rabbitmq", "redis", "docker"]
    assert intern.skills == ["pytest"]
    assert p.person.email == "petr.vymyslov@example.test"
    assert p.person.full_name == "Пётр Вымыслов" and p.person.city == "Казань"
    # стажировка 2020-06…2021-02 = 8 мес × 0.5, работа 2021-03…2026-10 = 67 мес
    assert p.total_experience_months == 4 + 67
    assert p.education[0].year_end == 2021


def test_assemble_rejects_invented_name_and_works_without_taxonomy():
    out = copy.deepcopy(LLM_OUT)
    out["full_name"] = "Иван Несуществующий"
    p = _assemble(out, index=None)
    assert p.person.full_name is None
    assert p.skills == [] and "taxonomy_unavailable" in p.warnings
    assert "Python" in p.unknown_skills and "Kubernetes" not in p.unknown_skills


# --- LLM: схема и повтор ---


async def test_llm_retry_once_then_fail():
    bad = {"skills": [{"name": "Python"}]}  # нет quote
    llm = FakeLlm(bad, LLM_OUT)
    got = await resume_parser.extract_llm(llm, TEXT)
    assert len(got.skills) == len(LLM_OUT["skills"])
    assert len(llm.calls) == 2
    assert "проверку схемы" in llm.calls[1]["messages"][-1]["content"]
    assert llm.calls[0]["schema"]["title"] == "LlmResume"

    with pytest.raises(LlmError) as e:
        await resume_parser.extract_llm(FakeLlm(bad, bad), TEXT)
    assert e.value.code == "llm_invalid"


def test_prompt_contains_text_and_rules():
    system, user = resume_parser.build_messages("Резюме кандидата X")
    assert "quote" in system["content"] and "дословный" in system["content"]
    assert "Резюме кандидата X" in user["content"]


async def test_openai_client_request_and_errors():
    seen = {}

    def ok(request: httpx.Request) -> httpx.Response:
        seen["body"] = json.loads(request.content)
        seen["auth"] = request.headers.get("authorization")
        msg = {"content": json.dumps({"a": 1})}
        return httpx.Response(200, json={"choices": [{"message": msg, "finish_reason": "stop"}]})

    llm = OpenAiLlm("http://llm/v1", "k", "m", transport=httpx.MockTransport(ok))
    assert await llm.complete_json([], {"type": "object"}, schema_name="resume") == {"a": 1}
    assert seen["auth"] == "Bearer k"
    assert seen["body"]["response_format"]["type"] == "json_schema"
    assert seen["body"]["response_format"]["json_schema"]["schema"] == {"type": "object"}
    assert seen["body"]["chat_template_kwargs"] == {"enable_thinking": False}

    def boom(request):
        raise httpx.ReadTimeout("slow")

    cases = [
        (httpx.MockTransport(boom), "llm_timeout"),
        (httpx.MockTransport(lambda r: httpx.Response(503)), "llm_unavailable"),
        (
            httpx.MockTransport(
                lambda r: httpx.Response(200, json={"choices": [{"message": {"content": "{"}}]})
            ),
            "llm_invalid",
        ),
    ]
    for transport, code in cases:
        with pytest.raises(LlmError) as e:
            await OpenAiLlm("http://llm/v1", "", "m", transport=transport).complete_json(
                [], {}, schema_name="x"
            )
        assert e.value.code == code


# --- обработчик команды ---


def _taxonomy() -> TaxonomyProvider:
    prov = TaxonomyProvider("", "")
    prov.set_snapshot(
        {"version": 3, "skills": yaml.safe_load((TAX / "skills.yaml").read_text())["skills"]}
    )
    return prov


def _command(object_key: str = "resumes/r.pdf"):
    p = cmd.MlResumeParse(
        resume_id=uuid4(),
        candidate_id=uuid4(),
        object_key=object_key,
        mime="application/pdf",
        sha256="0" * 64,
        taxonomy_version=3,
    )
    env, _ = new_command("ml.resume.parse", "core", p)
    return env, p


async def test_handler_completed_progress_and_redelivery():
    pdf = _pdf(TEXT)

    async def load(key: str) -> bytes:
        return pdf

    stages: list[tuple[str, int]] = []

    async def progress(stage, pct, message):
        stages.append((stage, pct))

    env, p = _command()
    kw = {"load": load, "taxonomy": _taxonomy(), "model": "qwen3-vl-4b", "today": TODAY}
    first = await resume_parse_llm(env, p, llm=FakeLlm(LLM_OUT), progress=progress, **kw)
    again = await resume_parse_llm(env, p, llm=FakeLlm(LLM_OUT), **kw)
    assert [e.type for e in first] == ["ml.resume.completed"]
    assert first[0].message_id == again[0].message_id  # inbox core отбросит дубль
    assert first[0].payload == again[0].payload
    parsed = first[0].payload["parsed"]
    assert parsed["parser_version"] == "resume-v2:qwen3-vl-4b"
    assert {s["slug"] for s in parsed["skills"]} >= {"python", "postgresql", "docker"}
    pcts = [x for _, x in stages]
    assert pcts == sorted(pcts) and stages[-1] == ("done", 100)


@pytest.mark.parametrize(
    ("data", "answer", "code"),
    [
        (b"not a file", LLM_OUT, "bad_file"),
        (None, LlmError("llm_timeout"), "llm_timeout"),
        ("missing", LLM_OUT, "file_unavailable"),
    ],
)
async def test_handler_failures(data, answer, code):
    pdf = _pdf(TEXT)

    async def load(key: str) -> bytes:
        if data == "missing":
            raise FileNotFoundError(key)
        return pdf if data is None else data

    env, p = _command()
    out = await resume_parse_llm(
        env, p, load=load, llm=FakeLlm(answer), taxonomy=_taxonomy(), model="m", today=TODAY
    )
    assert [e.type for e in out] == ["ml.resume.failed"]
    assert out[0].payload["error"] == code


# --- воркер с ML_FAKE=false ---


async def test_real_resume_parse_publishes_progress_and_result():
    """ML_FAKE=false: ml.resume.parse идёт через LLM, прогресс — в ui.user.<id>."""
    pdf = _pdf(TEXT)

    class Storage:
        async def get(self, key: str) -> bytes:
            return pdf

    w = MlWorker(
        "llm",
        Settings(ml_fake=False, llm_base_url=""),
        llm=FakeLlm(LLM_OUT),
        storage=Storage(),
        taxonomy=_taxonomy(),
    )
    events: list[dict] = []
    ui: list[dict] = []

    @w.broker.subscriber(durable_queue("t.ml.resume", "ml.resume.*"), EVENTS_EXCHANGE)
    async def sink(msg: dict):
        events.append(msg)

    @w.broker.subscriber(RabbitQueue("t.ui", routing_key="ui.user.*"), UI_EXCHANGE)
    async def ui_sink(msg: dict):
        ui.append(msg)

    p = cmd.MlResumeParse(
        resume_id=uuid4(),
        candidate_id=uuid4(),
        object_key="resumes/r.pdf",
        mime="application/pdf",
        sha256="0" * 64,
        taxonomy_version=3,
    )
    env, _ = new_command("ml.resume.parse", "core", p)
    async with TestRabbitBroker(w.broker):
        await w.consumer.dispatch(env)
    assert [m["type"] for m in events] == ["ml.resume.completed"]
    assert events[0]["payload"]["parsed"]["skills"]
    assert ui and all(m["event"] == "job.progress" for m in ui)
    assert ui[0]["user_id"] == str(p.candidate_id)
    assert ui[-1]["data"]["pct"] == 100 and ui[-1]["data"]["job_id"] == str(p.resume_id)


# --- заявленное в резюме (E20-5) ---

DECL_TEXT = (
    "Пётр Вымыслов\n"
    "Middle backend-разработчик\n"
    "Коммуникабельный, беру на себя наставничество джунов.\n"
    "Python, FastAPI, PostgreSQL\n"
)


def _decl_assemble(extra: dict, text: str = DECL_TEXT):
    out = resume_parser.LlmResume.model_validate(
        {"skills": [{"name": "Python", "quote": "Python, FastAPI, PostgreSQL"}], **extra}
    )
    return resume_parser.assemble(text, out, INDEX, version="resume-v2:test", today=TODAY)


def test_declared_fields_extracted_with_quotes():
    p = _decl_assemble(
        {
            "soft_skills": [
                {"name": "Коммуникабельность", "quote": "Коммуникабельный, беру на себя"},
                {"name": "Наставничество", "quote": "беру на себя наставничество джунов"},
            ],
            "declared_grade": {"grade": "middle", "quote": "Middle backend-разработчик"},
            "roles": [{"role": "backend", "quote": "Middle backend-разработчик"}],
        }
    )
    assert [s.name for s in p.soft_skills] == ["Коммуникабельность", "Наставничество"]
    assert all(s.quote for s in p.soft_skills)
    assert p.declared_grade is not None and p.declared_grade.grade == "middle"
    assert [r.role for r in p.declared_roles] == ["backend"]
    assert not any(w.endswith("_dropped") or "_dropped:" in w for w in p.warnings)


def test_declared_fields_dropped_without_quote():
    p = _decl_assemble(
        {
            "soft_skills": [
                {
                    "name": "Лидерство",
                    "quote": "Руководил командой из десяти человек",
                },  # нет в тексте
                {
                    "name": "Стрессоустойчивость",
                    "quote": "Коммуникабельный, беру на себя",
                },  # не в цитате
            ],
            "declared_grade": {"grade": "senior", "quote": "Senior-разработчик с 10 годами опыта"},
            "roles": [
                {"role": "devops", "quote": "Middle backend-разработчик"},  # роль не читается
                {"role": "frontend", "quote": "Фронтенд-разработчик на React"},  # нет в тексте
            ],
        }
    )
    assert p.soft_skills == [] and p.declared_grade is None and p.declared_roles == []
    for w in (
        "soft_dropped:Лидерство",
        "soft_dropped:Стрессоустойчивость",
        "grade_dropped:senior",
        "role_dropped:devops",
        "role_dropped:frontend",
    ):
        assert w in p.warnings


def test_declared_grade_must_match_quote():
    # грейд в цитате есть, но модель назвала другой
    p = _decl_assemble(
        {"declared_grade": {"grade": "senior", "quote": "Middle backend-разработчик"}}
    )
    assert p.declared_grade is None and "grade_dropped:senior" in p.warnings


def test_declared_does_not_change_skills_or_experience():
    """Заявленное не влияет ни на навыки, ни на стаж (категорию даёт только тест)."""
    base = _decl_assemble({})
    withd = _decl_assemble(
        {
            "soft_skills": [
                {"name": "Коммуникабельность", "quote": "Коммуникабельный, беру на себя"}
            ],
            "declared_grade": {"grade": "middle", "quote": "Middle backend-разработчик"},
            "roles": [{"role": "backend", "quote": "Middle backend-разработчик"}],
        }
    )
    assert withd.skills == base.skills and withd.experience == base.experience
    assert withd.total_experience_months == base.total_experience_months
    assert withd.unknown_skills == base.unknown_skills
    assert base.soft_skills == [] and base.declared_grade is None and base.declared_roles == []


def test_declared_schema_rejects_unknown_grade_and_role():
    with pytest.raises(ValidationError):
        resume_parser.LlmResume.model_validate({"declared_grade": {"grade": "guru", "quote": "x"}})
    with pytest.raises(ValidationError):
        resume_parser.LlmResume.model_validate({"roles": [{"role": "wizard", "quote": "x"}]})


def test_rules_mode_has_no_declared_fields():
    out = resume_parser.extract_rules(TEXT, INDEX)
    assert out.soft_skills == [] and out.declared_grade is None and out.roles == []


# --- E43-8: запасной путь, когда LLM вернул мало навыков ---


async def _parse_with(answer, text=TEXT):
    return await resume_parser.parse_resume(
        _pdf(text), llm=FakeLlm(answer), index=INDEX, model="test", today=TODAY
    )


async def test_fallback_rules_when_llm_returns_zero_skills():
    parsed = await _parse_with({**LLM_OUT, "skills": []})
    slugs = {s.slug for s in parsed.skills}
    assert {"python", "fastapi", "postgresql", "docker"} <= slugs
    assert all(s.source == "rules" and s.quote.strip() for s in parsed.skills)
    assert any(w.startswith("llm_fallback_rules:") for w in parsed.warnings)
    assert parsed.parser_version == "resume-v2:test"  # разбор остаётся LLM-овым, опыт от LLM
    assert parsed.experience


async def test_fallback_keeps_llm_skills_and_adds_only_new_with_quote():
    one = {**LLM_OUT, "skills": [LLM_OUT["skills"][0]]}
    long_text = TEXT + "\n" + "Дополнительная информация о проектах. " * 40
    parsed = await _parse_with(one, long_text)
    by = {s.slug: s for s in parsed.skills}
    assert by[one["skills"][0]["name"].lower()].source == "llm"
    assert any(s.source == "rules" for s in parsed.skills)
    assert len(by) == len(parsed.skills)  # без дублей
    text_norm = resume_parser.norm_text(long_text)
    assert all(resume_parser.quote_in_text(s.quote, text_norm) for s in parsed.skills)


async def test_no_fallback_when_llm_enough_or_short_text():
    full = await _parse_with(LLM_OUT)
    assert all(s.source == "llm" for s in full.skills)
    assert not any(w.startswith("llm_fallback_rules") for w in full.warnings)
    # 1–2 навыка при коротком тексте (< 1000 знаков) — не повод для запасного пути
    short = await _parse_with({**LLM_OUT, "skills": LLM_OUT["skills"][:1]})
    assert len(short.skills) == 1 and short.skills[0].source == "llm"


def test_needs_rules_fallback_thresholds():
    f = resume_parser.needs_rules_fallback
    assert f(0, 200) and f(0, 5000)
    assert f(2, 1001) and not f(2, 1000)
    assert not f(3, 5000)
    assert not f(0, 0)


# --- E44-2: цитата без названия навыка — ищем название рядом, а не теряем навык ---

FRONT_TEXT = (FIX / "resume_frontend.txt").read_text()
# (навык, цитата из резюме без его названия): так отвечает слабая модель
FRONT_REAL = [
    ("JavaScript", "Разрабатываю личный кабинет клиента для сервиса онлайн-записи."),
    ("HTML", "Отвечал за адаптивную вёрстку под телефоны и планшеты."),
    ("Redux", "Перевёл сложные формы на новую архитектуру."),
    ("Vue3", "Настроил сборку проекта и проверку стиля кода в CI."),
    ("TypeScript", "Работаю в связке с дизайнерами и бэкендерами, согласую контракты API."),
    ("CSS", "Писал компоненты для таблиц, фильтров и графиков."),
    ("React", "Поддерживал внутренний портал сотрудников и административную панель."),
    ("Webpack", "Участвую в митапах по фронтенду, веду заметки о новых возможностях браузеров."),
    ("Git", "Освоил основы работы с системой контроля версий и командной разработки."),
    ("npm", "Покрываю компоненты тестами, пишу документацию к интерфейсным библиотекам команды."),
    ("Sass", "Библиотека компонентов: кнопки, поля ввода, модальные окна, таблицы с сортировкой."),
    ("Figma", "Верстал лендинги по макетам из Figma, подключал формы обратной связи."),
]
FRONT_FAKE = [
    ("Angular", "Перевёл сложные формы на новую архитектуру."),
    ("Java", "Пишу на JavaScript и TypeScript, строю одностраничные приложения на Vue3 и React."),
    ("Kubernetes", "Разворачивал кластеры и настраивал автоскейлинг."),
]


def _front_assemble(pairs):
    out = resume_parser.LlmResume(
        skills=[resume_parser.LlmSkill(name=n, quote=q) for n, q in pairs]
    )
    return resume_parser.assemble(FRONT_TEXT, out, INDEX, version="t", today=TODAY)


def test_recover_quote_real_skills_not_lost():
    p = _front_assemble(FRONT_REAL + FRONT_FAKE)
    got = {s.slug for s in p.skills}
    want = {"javascript", "html", "redux", "vue", "typescript", "css", "react", "git", "npm"}
    assert want <= got
    lost = [n for n, _ in FRONT_REAL if f"skill_not_in_quote:{n}" in p.warnings]
    assert len(lost) / len(FRONT_REAL) < 0.10, lost
    text_norm = resume_parser.norm_text(FRONT_TEXT)
    aliases = resume_parser._aliases_by_slug(INDEX)
    for s in p.skills:  # цитата — из текста, и в ней читается навык (название или синоним)
        assert resume_parser.quote_in_text(s.quote, text_norm)
        norm_quote = resume_parser.norm_text(s.quote)
        assert resume_parser.mention_in_quote(s.raw, norm_quote, aliases[s.slug])
    by = {s.slug: s for s in p.skills}
    assert "JavaScript" in by["javascript"].quote  # цитата со строкой, где есть название


def test_recover_quote_keeps_dropping_skills_without_name():
    p = _front_assemble(FRONT_FAKE)
    assert p.skills == []
    assert "skill_not_in_quote:Angular" in p.warnings
    assert "skill_not_in_quote:Java" in p.warnings  # Java ≠ JavaScript даже при поиске по тексту
    assert "quote_not_found:Kubernetes" in p.warnings


def test_recover_quote_word_boundaries():
    text = "Пишу на C# и JavaScript.\nСтек: C, Go.\nДругая строка."
    assert resume_parser.recover_quote(text, "Другая строка.", "C", []) == "Стек: C, Go."
    assert resume_parser.recover_quote(text, "Другая строка.", "Java", []) is None
    assert resume_parser.recover_quote("Пишу на C# и Go.", "Пишу на Go.", "C", []) is None
